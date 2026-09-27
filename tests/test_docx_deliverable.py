from pathlib import Path
from types import SimpleNamespace

from app.state.workflow_state import WorkflowState
from app.workflow.graph import WorkflowGraph


def test_deliverable_accepts_docx_format(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "outputs").mkdir()

    state = WorkflowState(
        request_id="req_docx_test",
        user_query="Generate a report",
        requested_deliverable="docx",
        final_answer="Verified report answer",
        verification_status="passed",
        verification_results=[
            {
                "verification_status": "passed",
                "failures": [],
                "numeric_validation": [],
            }
        ],
    )

    graph = WorkflowGraph.__new__(WorkflowGraph)
    graph.telemetry = SimpleNamespace(
        summary=lambda: {}
    )

    result = graph._deliverable(state)

    assert len(result.generated_deliverables) == 1
    assert Path(result.generated_deliverables[0]) == Path(
        "outputs/req_docx_test_report.docx"
    )

    output_file = Path("outputs/req_docx_test_report.docx")

    assert output_file.exists()
    assert output_file.stat().st_size > 0
    assert output_file.read_bytes()[:4] == b"PK\x03\x04"
def test_report_deliverable_uses_sovara_structure(
    tmp_path,
    monkeypatch,
):
    from pathlib import Path
    from types import SimpleNamespace

    from docx import Document

    from app.state.workflow_state import WorkflowState
    from app.workflow.graph import WorkflowGraph

    monkeypatch.chdir(tmp_path)
    (tmp_path / "outputs").mkdir()

    telemetry = {
        "llm_calls": 3,
        "external_api_calls": 0,
        "network_calls": 0,
        "cloud_uploads": 0,
        "models_used": ["phi4-mini:latest"],
        "tools_used": [
            "ReasoningNode",
            "SynthesisNode",
        ],
        "files_processed": 1,
        "sandbox_executions": 0,
        "sandbox_successes": 0,
        "sandbox_failures": 0,
        "local_inference": True,
        "processing_location": "local",
        "no_external_calls": True,
    }

    state = WorkflowState(
        request_id="req_report_structure",
        user_query="Generate a structured report",
        requested_deliverable="report",
        final_answer=(
            "Executive Summary:\n"
            "Revenue increased based on the supplied analysis.\n\n"
            "Key Findings:\n"
            "- Revenue increased by the measured amount.\n\n"
            "Detailed Analysis:\n"
            "The supplied analysis contains the supporting values.\n\n"
            "Conclusion:\n"
            "The analysis confirms the observed change."
        ),
        synthesis_result={
            "answer": "Structured report answer",
        },
        retrieved_evidence=[
            {
                "evidence_id": "ev-001",
                "source_filename": "source.pdf",
                "page_number": 4,
                "chunk_id": "chunk-7",
                "citation_id": "cite-001",
                "content": "Important source evidence.",
            }
        ],
        document_results=[
            {
                "evidence_id": "doc-001",
                "source_filename": "source.pdf",
            }
        ],
        data_results=[
            {
                "evidence_id": "data-001",
                "result": {
                    "total": 550,
                },
            }
        ],
        vision_results=[
            {
                "evidence_id": "vision-001",
                "result": {
                    "observations": [],
                },
            }
        ],
        verification_status="passed",
        verification_results=[
            {
                "verification_status": "passed",
                "failures": [],
                "numeric_validation": [],
            }
        ],
        execution_telemetry=telemetry,
        generated_deliverables=[],
    )

    graph = WorkflowGraph.__new__(WorkflowGraph)
    graph.telemetry = SimpleNamespace(
        summary=lambda: telemetry
    )

    result = graph._deliverable(state)

    assert len(result.generated_deliverables) == 1

    output_file = Path(
        "outputs/req_report_structure_report.docx"
    )

    assert output_file.exists()

    document = Document(str(output_file))

    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading")
    ]

    assert "SOVARA ANALYSIS REPORT" in headings
    assert "1 Executive Summary" in headings
    assert "2 Key Findings" in headings
    assert "3 Detailed Analysis" in headings
    assert "4 Supporting Evidence" in headings
    assert "5 Verification" in headings
    assert "6 Conclusion" in headings
    assert "Appendix" in headings
    assert "A. Evidence Registry" in headings
    assert "B. Execution & Sovereignty Telemetry" in headings
    assert "C. Generated Artifacts" in headings

    assert "Answer" not in headings
    assert "Evidence" not in headings
    assert "Documents" not in headings
    assert "Data" not in headings
    assert "Vision" not in headings
    assert "Verification Status" not in headings

    document_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "Important source evidence." in document_text
    assert "ev-001" in document_text
    assert "source.pdf" in document_text
    assert "llm_calls" in document_text
    assert "local_inference" in document_text
    assert "no_external_calls" in document_text


def test_json_deliverable_uses_machine_readable_sovara_structure(
    tmp_path,
    monkeypatch,
):
    import json
    from pathlib import Path
    from types import SimpleNamespace

    from app.state.workflow_state import WorkflowState
    from app.workflow.graph import WorkflowGraph

    monkeypatch.chdir(tmp_path)
    (tmp_path / "outputs").mkdir()

    telemetry = {
        "llm_calls": 2,
        "external_api_calls": 0,
        "network_calls": 0,
        "cloud_uploads": 0,
        "local_inference": True,
        "processing_location": "local",
        "no_external_calls": True,
    }

    state = WorkflowState(
        request_id="req_json_structure",
        user_query="Generate JSON report",
        requested_deliverable="json",
        final_answer="Analysis completed.",
        retrieved_evidence=[
            {
                "evidence_id": "ev-json-001",
                "source_filename": "source.csv",
                "page_number": 2,
                "chunk_id": "chunk-4",
                "content": "Source evidence.",
            }
        ],
        document_results=[
            {
                "document_id": "doc-001",
            }
        ],
        data_results=[
            {
                "evidence_id": "data-001",
                "result": {
                    "sum": 550,
                },
            }
        ],
        vision_results=[],
        verification_status="passed",
        verification_results=[
            {
                "verification_status": "passed",
                "failures": [],
                "numeric_validation": [],
            }
        ],
        execution_telemetry=telemetry,
        generated_deliverables=[],
    )

    graph = WorkflowGraph.__new__(WorkflowGraph)
    graph.telemetry = SimpleNamespace(
        summary=lambda: telemetry
    )

    result = graph._deliverable(state)

    assert len(result.generated_deliverables) == 1

    payload = json.loads(
        Path(
            result.generated_deliverables[0]
        ).read_text(encoding="utf-8")
    )

    assert payload["title"] == "SOVARA ANALYSIS REPORT"

    assert set(payload["sections"]) == {
        "executive_summary",
        "key_findings",
        "detailed_analysis",
        "supporting_evidence",
        "verification",
        "conclusion",
    }

    assert set(payload["appendix"]) == {
        "evidence_registry",
        "execution_telemetry",
        "generated_artifacts",
    }

    assert "answer" not in payload["sections"]
    assert "evidence" not in payload["sections"]
    assert "documents" not in payload["sections"]
    assert "data" not in payload["sections"]
    assert "vision" not in payload["sections"]
    assert "verification_status" not in payload["sections"]

    registry = payload["appendix"]["evidence_registry"]

    assert (
        registry["retrieved_evidence"][0]["evidence_id"]
        == "ev-json-001"
    )
    assert (
        registry["retrieved_evidence"][0]["page_number"]
        == 2
    )
    assert (
        registry["retrieved_evidence"][0]["chunk_id"]
        == "chunk-4"
    )
    assert (
        registry["document_results"][0]["document_id"]
        == "doc-001"
    )
    assert (
        registry["data_results"][0]["evidence_id"]
        == "data-001"
    )

    assert payload["sections"]["verification"]["status"] == "passed"

    assert (
        payload["sections"]["verification"]["results"][0][
            "verification_status"
        ]
        == "passed"
    )

    assert (
        payload["appendix"]["execution_telemetry"][
            "no_external_calls"
        ]
        is True
    )