from __future__ import annotations

from pathlib import Path
from docx import Document


class DOCXParser:

    @staticmethod
    def extract_text(docx_path: str | Path) -> dict:
        path = Path(docx_path)
        document = Document(path)

        paragraphs = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        tables = []

        for table in document.tables:
            rows = []

            for row in table.rows:
                rows.append([
                    cell.text.strip()
                    for cell in row.cells
                ])

            if rows:
                tables.append(rows)

        return {
            "filename": path.name,
            "paragraphs": paragraphs,
            "tables": tables,
            "text": "\n".join(paragraphs),
        }

    @staticmethod
    def chunk_text(text: str, chunk_size: int = 1000) -> list[dict]:
        text = text.strip()

        if not text:
            return []

        chunks = []

        for index in range(0, len(text), chunk_size):
            chunks.append({
                "chunk_id": f"chunk_{index // chunk_size + 1:05d}",
                "page_number": None,
                "text": text[index:index + chunk_size],
                "ocr_used": False,
            })

        return chunks