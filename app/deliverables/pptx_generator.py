from __future__ import annotations

from pathlib import Path

from pptx import Presentation


class PptxGenerator:
    def generate(self, slides: list[str], output_path: str | Path) -> str:
        prs = Presentation()
        for slide_text in slides:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = "Findings"
            slide.shapes.placeholders[1].text = slide_text
        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out))
        return str(out)
