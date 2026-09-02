from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx import Document


class DocxGenerator:
    @staticmethod
    def _to_text(value: Any) -> str:
        if value is None:
            return ""

        if isinstance(value, str):
            return value

        if isinstance(value, (dict, list)):
            return json.dumps(
                value,
                indent=2,
                ensure_ascii=False,
                default=str,
            )

        return str(value)

    def generate(self, report: dict, output_path: str | Path) -> str:
        document = Document()

        document.add_heading(
            report.get("title", "Analysis Report"),
            level=1,
        )

        for section_title, section_content in report.get("sections", {}).items():
            document.add_heading(
                str(section_title).replace("_", " ").title(),
                level=2,
            )

            document.add_paragraph(
                self._to_text(section_content)
            )

        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)

        document.save(str(out))

        return str(out)